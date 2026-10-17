"""
map_benchmarks.py -- decide which index each fund is compared against.
----------------------------------------------------------------------
Reads the CRISIL "Fund Performance" spreadsheets, which name each
scheme's own stated benchmark, and writes mf_benchmark_map.

USAGE
    python map_benchmarks.py "benchmark mutual fund references"
    python map_benchmarks.py "..." --dry-run
    python map_benchmarks.py "..." --no-category      # per-scheme only
    python map_benchmarks.py "..." --include-unscored # cover unscored funds

WHAT IT WRITES
    EXACT     we hold the index the fund actually states -- or, for an index
              fund / ETF, the index its own NAME says it tracks (see the
              name-based step below; the note says which)
    PROXY     we hold a near-equivalent (BSE 500 shown against Nifty 500)
    CATEGORY  no stated benchmark found, but the fund sits in a category
              whose AMFI Tier-1 benchmark we hold
    NONE      nothing defensible; the site should show no comparison

WHY 'NONE' IS A RESULT AND NOT A FAILURE
A sectoral fund benchmarked to Nifty Financial Services has no honest
stand-in among broad-market indices. Showing one would produce a number
that looks authoritative and means nothing. Better to show nothing and
keep the fund on the review list until the right index is loaded.

WHY 'CATEGORY' EXISTS
The per-scheme path only works for funds that appear in a reference file
we happen to have downloaded. Whole categories were sitting at zero
resolved simply because their file was never downloaded. A Mid Cap fund
is benchmarked to Nifty Midcap 150 whether or not CRISIL's spreadsheet
is on disk, so the category default is defensible on its own terms --
it is what AMFI's Tier-1 category benchmark list says.

It is deliberately weaker than a stated match: confidence never exceeds
0.70, the note says plainly that this is a category benchmark rather than
the scheme's own, and an EXACT or PROXY match always wins over it. The
fund page must disclose the difference.

Categories deliberately EXCLUDED from the fallback are listed in the
CATEGORY_EXCLUDED note below.

It also writes benchmark_review.csv -- every fund still needing a human
decision, so the gap is a worklist rather than a surprise.

ON THE SOURCE FILES: they are used here only to read each scheme's
stated benchmark. The returns shown on the site are computed from AMFI
NAV and NSE index values -- CRISIL's own return figures are not loaded
and should not be republished.
"""

import os
import re
import csv
import sys
import glob
import argparse

import psycopg
import openpyxl
from dotenv import load_dotenv

load_dotenv()
DB = os.getenv("FINCHAYA_DB")
if not DB:
    sys.exit("FINCHAYA_DB is not set. Check that .env exists in this folder.")

REVIEW_CSV = "benchmark_review.csv"


# ---------------------------------------------------------------------
# PROXIES. Only where the two indices track substantially the same
# market -- BSE 500 and Nifty 500 differ by a few basis points a year.
# The note travels with the mapping so the page can disclose it.
#
# Deliberately NOT here: any sector or theme index. Financial Services
# is not "close enough" to Nifty 500 in any period that matters.
# ---------------------------------------------------------------------
# Keys are in norm_bm() form: lowercased, punctuation stripped, and the
# TRI marker removed -- so "BSE 500 TRI" arrives here as "bse 500".
PROXIES = {
    "bse 500":              ("NIFTY 500",             0.90),
    "bse 200":              ("NIFTY 500",             0.80),
    "nifty 200":            ("NIFTY 500",             0.85),
    "bse 100":              ("NIFTY 100",             0.90),
    "bse sensex":           ("NIFTY 50",              0.85),
    "bse midcap 150":       ("NIFTY MIDCAP 150",      0.90),
    "bse 250 smallcap":     ("NIFTY SMALLCAP 250",    0.90),
    "bse 250 large midcap": ("NIFTY LARGEMIDCAP 250", 0.90),
}


# ---------------------------------------------------------------------
# CATEGORY FALLBACK. AMFI's Tier-1 benchmark for the category, used only
# when no stated benchmark could be found for the scheme.
#
# Keys are in norm() form -- lowercase, alphanumeric only -- so both
# "Mid Cap Fund" and "Midcap Fund" arrive here as "midcapfund", and
# "ELSS- Tax Saver Fund" as "elsstaxsaverfund".
#
# Confidence is graded by how well one broad index actually represents
# the category. Large/Mid/Small/LargeMid map to their own index and
# score highest. Flexi Cap and ELSS have Nifty 500 as their genuine
# Tier-1 benchmark. Focused, Value and Contra are mandated to invest
# across the market but with a concentrated or style-tilted book, so
# Nifty 500 is the right family with a wider dispersion -- scored lower.
# ---------------------------------------------------------------------
CATEGORY_BENCHMARK = {
    "largecapfund":        ("NIFTY 100",             0.70),
    "midcapfund":          ("NIFTY MIDCAP 150",      0.70),
    "smallcapfund":        ("NIFTY SMALLCAP 250",    0.70),
    "largeandmidcapfund":  ("NIFTY LARGEMIDCAP 250", 0.70),
    "flexicapfund":        ("NIFTY 500",             0.65),
    "elss":                ("NIFTY 500",             0.65),
    "elsstaxsaverfund":    ("NIFTY 500",             0.65),
    "focusedfund":         ("NIFTY 500",             0.60),
    "valuefund":           ("NIFTY 500",             0.60),
    "contrafund":          ("NIFTY 500",             0.60),
    # Held since Aug 2026. AMFI's Tier-1 for the category: weights large,
    # mid and small at 50:25:25 rather than by float, which is the whole
    # reason plain Nifty 500 was refused here before the series existed.
    "multicapfund":        ("NIFTY 500 MULTICAP 50:25:25", 0.70),
}

# Categories deliberately given NO fallback, and why:
#
#   Multi Cap Fund      -- NO LONGER EXCLUDED. It sat here while we did
#                          not hold Nifty 500 Multicap 50:25:25 TRI,
#                          because plain Nifty 500 measures a multi cap
#                          fund's mid and small exposure against the
#                          wrong mix -- a methodology difference, not a
#                          rounding error. The series was loaded in
#                          Aug 2026, so the category now has its real
#                          Tier-1 benchmark and maps like any other.
#   Dividend Yield Fund -- AMCs split between Nifty 500 and Nifty
#                          Dividend Opportunities 50. No single default
#                          is defensible across the category.
#   Sectoral / Thematic -- the whole point of NONE. See module docstring.
#   Index Fund / ETF    -- benchmark is the fund's own tracked index;
#                          belongs in the name-based fallback, not here.
CATEGORY_EXCLUDED = {
    "dividendyieldfund",
    "sectoralthematic", "thematicfund", "sectoralfund",
    "indexfund", "etfs", "etf",
}

CATEGORY_MAX_CONFIDENCE = 0.70


def norm(text):
    """Compare scheme names loosely enough to survive punctuation and
    spacing differences, strictly enough not to merge two funds.

    All non-alphanumerics are dropped, spaces included. AMFI and CRISIL
    disagree on the space in "Mid Cap" vs "Midcap" -- and disagree in
    BOTH directions, so a one-way rewrite would not fix it. Compare
    without spaces and the question disappears.

    Dropping spaces raises the theoretical risk of two distinct funds
    colliding on one key; read_references() logs any collision rather
    than resolving it silently.
    """
    t = (text or "").lower().replace("&", " and ")
    return re.sub(r"[^a-z0-9]+", "", t)


def norm_bm(text):
    """Normalise an index name for comparison.

    The NSE files call it 'NIFTY 500'; funds state 'Nifty 500 TRI'. Both
    mean the same index, so the TRI marker is stripped before matching --
    otherwise every exact match silently becomes a miss.

    Spaces are KEPT here, unlike norm(): PROXIES keys are written in
    spaced form and index names are short enough that spacing is stable.
    """
    t = (text or "").lower().replace("&", " and ")
    t = re.sub(r"[^a-z0-9 ]+", " ", t)
    t = re.sub(r"\s+", " ", t).strip()
    t = re.sub(r"\b(total returns? index|tri|index)\b", " ", t)
    return re.sub(r"\s+", " ", t).strip()


def resolve_index(have, target):
    """benchmark_id for an index name, tolerating the one space NSE and
    CRISIL disagree on: "NIFTY500 MULTICAP 50:25:25" against
    "Nifty 500 Multicap 50:25:25 TRI". Same disagreement norm() handles
    for scheme names, and it runs in both directions here too.

    Exact norm_bm() match first, so nothing already working changes.
    Only on a miss do we compare with spaces removed, and only when
    EXACTLY ONE index matches -- two indices collapsing onto one key is
    left unmapped rather than guessed at, because picking the wrong
    index is worse than showing no benchmark.
    """
    key = norm_bm(target)
    if key in have:
        return have[key]
    squashed = key.replace(" ", "")
    if not squashed:
        return None
    hits = {bid for k, bid in have.items() if k.replace(" ", "") == squashed}
    return hits.pop() if len(hits) == 1 else None


# ---------------------------------------------------------------------
# NAME-BASED MATCH FOR INDEX FUNDS AND ETFs
#
# An index fund's benchmark is the index it tracks, and the fund says so in
# its name: "Axis Nifty Next 50 Index Fund" tracks the Nifty Next 50. The
# CRISIL sheets rarely list these funds, so without this step they sit at
# NONE however many indices we hold (the category fallback skips them on
# purpose: there is no category benchmark for "Index Fund").
#
# THE RISK IS A WRONG MATCH, NOT A MISSED ONE. "Nifty 50 Equal Weight Index
# Fund" does not track the Nifty 50, and a Nifty 50 comparison would be
# confidently wrong. So the rule is strict and fails toward NONE:
#
#   * the name must say Index / ETF / BeES / FoF;
#   * a held index must appear in it as WHOLE WORDS, longest first
#     ("nifty 500" is not found inside "nifty 500 multicap 50 25 25");
#   * after the index only filler words may follow (Index, Fund, ETF, Direct,
#     Growth, ...). Anything else -- Equal Weight, Value 20, Momentum 30,
#     Quality, a second index -- means it is a DIFFERENT index, so no match;
#   * before it only the AMC's own name and filler (at most two unknown words,
#     and never Equal / Value / Momentum / ... ), so a brand quirk like an old
#     AMC name does not block a match but a style qualifier does;
#   * two different held indices in one name: no match.
#
# Whatever fails is left for a human and the review file says what the name
# looks like it tracks, which doubles as the list of indices worth downloading.
# ---------------------------------------------------------------------
PASSIVE_MARKERS = {"index", "etf", "etfs", "bees", "fof"}

# Words that say nothing about WHICH index a fund tracks.
NAME_NOISE = {
    "index", "indices", "fund", "funds", "etf", "etfs", "exchange", "traded", "tri", "total",
    "return", "returns", "direct", "regular", "plan", "growth", "idcw", "option", "options",
    "dividend", "payout", "reinvestment", "of", "the", "fof", "mf", "mutual", "scheme", "bees",
    "units", "unit", "open", "ended", "passive", "passively", "managed",
}

# Words that mean "a variation of the index", never allowed in front of it.
NAME_QUALIFIERS = {
    "equal", "weight", "weighted", "value", "momentum", "quality", "alpha", "low", "volatility",
    "esg", "sectoral", "sector", "enhanced", "smart", "factor", "factors", "high", "beta", "ex",
}

MAX_UNKNOWN_PREFIX = 2
NAME_CONFIDENCE = 0.95      # the name states the index outright
ALIAS_CONFIDENCE = 0.90     # a well-known short name ("Nifty Index Fund", "Junior BeES")

# Short names that are unambiguous in practice. Only used if the index is held.
NAME_ALIASES = {
    ("nifty",):          "NIFTY 50",       # "UTI Nifty Index Fund", "ICICI Prudential Nifty ETF"
    ("nifty", "bees"):   "NIFTY 50",
    ("junior", "bees"):  "NIFTY NEXT 50",
    ("bank", "bees"):    "NIFTY BANK",
}


def name_tokens(text):
    """Lowercase words of a fund or index name, with the spellings AMCs disagree on made the same:
    'Nifty500' / 'Nifty 500', 'Mid Cap' / 'Midcap', 'Large Midcap' / 'LargeMidcap', '&' / 'and'."""
    t = (text or "").lower().replace("&", " and ")
    t = re.sub(r"\blarge\s*(?:and\s*)?mid\s*cap\b", "largemidcap", t)
    t = re.sub(r"\bmid\s+cap\b", "midcap", t)
    t = re.sub(r"\bsmall\s+cap\b", "smallcap", t)
    t = re.sub(r"\bmicro\s+cap\b", "microcap", t)
    t = re.sub(r"\bmulti\s+cap\b", "multicap", t)
    t = re.sub(r"\b(nifty|midcap|smallcap|largemidcap|microcap|multicap|next|sensex)(?=\d)", r"\1 ", t)
    return re.sub(r"[^a-z0-9]+", " ", t).split()


def index_phrases(index_names):
    """({word tuple: index name} for every index we hold, {word tuple: index name} for the short names)."""
    full = {tuple(name_tokens(n)): n for n in index_names}
    held = set(index_names)
    return full, {ph: idx for ph, idx in NAME_ALIASES.items() if idx in held}


def index_from_name(name, amc_name, phrases, aliases):
    """(index name, 'name' | 'alias', words matched) for the held index an index fund / ETF names,
    or (None, reason, None). See the block comment above for the rules."""
    toks = name_tokens(name)
    if not (PASSIVE_MARKERS & set(toks)):
        return None, "not an index fund or ETF by name", None
    hits = []
    for table, is_alias in ((phrases, False), (aliases, True)):
        for ph, idx in table.items():
            n = len(ph)
            for i in range(len(toks) - n + 1):
                if tuple(toks[i:i + n]) == ph:
                    hits.append((n, i, idx, is_alias))
    if not hits:
        return None, "names no index we hold", None
    hits.sort(key=lambda h: (-h[0], h[1], h[3]))          # longest first, then leftmost, full names before aliases
    spans = []
    for n, i, idx, is_alias in hits:
        if not any(i < s_i + s_n and s_i < i + n for s_n, s_i, _, _ in spans):      # skip anything inside a longer hit
            spans.append((n, i, idx, is_alias))
    if len({s[2] for s in spans}) > 1:
        return None, "names more than one index: %s" % ", ".join(sorted({s[2] for s in spans})), None
    n, i, idx, is_alias = spans[0]
    prefix, suffix = toks[:i], toks[i + n:]
    after = [t for t in suffix if t not in NAME_NOISE]
    if after:
        return None, "the index is followed by '%s' (a variation or a different index)" % " ".join(after), None
    amc = set(name_tokens(amc_name))
    unknown = [t for t in prefix if t not in amc and t not in NAME_NOISE]
    if len(unknown) > MAX_UNKNOWN_PREFIX or any(t in NAME_QUALIFIERS for t in unknown):
        return None, "unrecognised words before the index: '%s'" % " ".join(unknown), None
    return idx, ("alias" if is_alias else "name"), " ".join(toks[i:i + n])


def tracked_guess(name):
    """What an index fund's name seems to track, for the review file: everything from the first index
    brand word on, minus filler. A hint for a human, never used to link anything."""
    toks = name_tokens(name)
    start = next((k for k, t in enumerate(toks) if t in ("nifty", "bse", "sensex", "crisil", "msci", "nasdaq", "s", "gold", "silver")), None)
    core = [t for t in (toks[start:] if start is not None else toks) if t not in NAME_NOISE]
    return " ".join(core)


def read_references(folder):
    """{normalised scheme name: (original name, stated benchmark)}"""
    paths = sorted(glob.glob(os.path.join(folder, "**", "*.xlsx"), recursive=True))
    if not paths:
        sys.exit("No .xlsx files found under %s" % folder)

    out = {}
    collisions = []
    for path in paths:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            started = False
            for row in ws.iter_rows(values_only=True):
                if not started:
                    # The real header sits several rows down, under a
                    # title block -- find it rather than assume a row.
                    if row and row[0] and str(row[0]).strip() == "Scheme Name":
                        started = True
                    continue
                if not row or not row[0]:
                    continue
                name = str(row[0]).strip()
                bm = str(row[1]).strip() if len(row) > 1 and row[1] else None
                if not (name and bm):
                    continue
                key = norm(name)
                prior = out.get(key)
                if prior and prior[0] != name:
                    # Two different printed names normalising to one key.
                    # Keep the first and surface it -- do not guess.
                    collisions.append((key, prior[0], name))
                    continue
                out[key] = (name, bm)

    print("Read %d files, %d schemes with a stated benchmark" % (len(paths), len(out)))
    if collisions:
        print("\n  WARNING: %d name collision(s) after normalisation." % len(collisions))
        print("  Kept the first; check these are genuinely the same fund:")
        for key, first, second in collisions[:20]:
            print("    %-40s  <-  %s | %s" % (key, first, second))
        if len(collisions) > 20:
            print("    ... and %d more" % (len(collisions) - 20))
        print()
    return out


def load_funds(conn, include_unscored):
    """Funds needing a benchmark, one representative scheme_code each.

    Default: funds carrying a current health score. mf_score holds one
    scheme_code per fund, which is what keeps mf_benchmark_map at one
    row per fund.

    With --include-unscored: also active equity funds with no score yet
    (typically because their AMC's holdings file has not been loaded).
    Those have no scheme_code chosen for them, so DIRECT/GROWTH is used
    -- the variant the fund page shows. Funds already covered by the
    scored set are not added twice.
    """
    with conn.cursor() as cur:
        cur.execute("""
            SELECT DISTINCT m.scheme_code, m.scheme_name,
                   COALESCE(m.sub_category, m.category) AS category
            FROM mf_scheme m
            JOIN mf_score s USING (scheme_code)
            WHERE s.as_of_date = (SELECT MAX(as_of_date) FROM mf_score)
            ORDER BY 2
        """)
        funds = cur.fetchall()
    print("Scored funds needing a benchmark: %d" % len(funds))

    if not include_unscored:
        print()
        return funds

    seen = {norm(name) for _, name, _ in funds}
    with conn.cursor() as cur:
        cur.execute("""
            SELECT DISTINCT ON (m.scheme_name)
                   m.scheme_code, m.scheme_name,
                   COALESCE(m.sub_category, m.category) AS category
            FROM mf_scheme m
            WHERE m.is_active
              AND m.category ILIKE '%Equity%'
              AND m.plan_type = 'DIRECT'
              AND m.option_type = 'GROWTH'
            ORDER BY m.scheme_name, m.scheme_code
        """)
        extra = [r for r in cur.fetchall() if norm(r[1]) not in seen]

    print("Unscored active equity funds added: %d" % len(extra))
    print("  (mapped on their DIRECT/GROWTH scheme code)\n")
    return funds + extra


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("folder", help="folder of CRISIL Fund Performance xlsx files")
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--no-category", action="store_true",
                    help="skip the category fallback pass")
    ap.add_argument("--include-unscored", action="store_true",
                    help="also map active equity funds that have no health score")
    ap.add_argument("--list-name-matches", action="store_true",
                    help="print every index fund / ETF linked by its name, grouped by index, to eyeball")
    args = ap.parse_args()

    refs = read_references(args.folder)

    with psycopg.connect(DB) as conn:
        with conn.cursor() as cur:
            cur.execute("SELECT index_name, benchmark_id FROM benchmark_master")
            index_ids = dict(cur.fetchall())                  # NSE index name -> benchmark_id
            have = {norm_bm(n): bid for n, bid in index_ids.items()}
            if not have:
                sys.exit("benchmark_master is empty -- run load_benchmarks.py first.")
            print("Indices loaded: %d\n" % len(have))

            # A mapping a human has confirmed is not re-derived. The old
            # UPDATE overwrote benchmark_id while leaving reviewed = true,
            # so a confirmed row could quietly start pointing elsewhere.
            cur.execute("""
                SELECT scheme_code FROM mf_benchmark_map WHERE reviewed
            """)
            reviewed = {r[0] for r in cur.fetchall()}
        if reviewed:
            print("Human-confirmed mappings left untouched: %d\n" % len(reviewed))

        funds = load_funds(conn, args.include_unscored)
        with conn.cursor() as cur:
            cur.execute("SELECT scheme_code, amc_name FROM mf_scheme WHERE scheme_code = ANY(%s)",
                        ([f[0] for f in funds],))
            amc_of = dict(cur.fetchall())
        phrases, aliases = index_phrases(list(index_ids))

        rows, review = [], []
        counts = {"EXACT": 0, "NAME": 0, "CONFLICT": 0, "PROXY": 0, "CATEGORY": 0,
                  "NONE_INDEX": 0, "NONE_NO_REF": 0, "SKIPPED": 0}
        cat_hits = {}
        name_hits = {}

        for code, name, category in funds:
            if code in reviewed:
                counts["SKIPPED"] += 1
                continue

            ref = refs.get(norm(name))
            stated = ref[1] if ref else None
            key = norm_bm(stated) if stated else None
            # What the fund's own NAME says it tracks (index funds and ETFs only).
            nm_idx, nm_kind, nm_detail = index_from_name(name, amc_of.get(code), phrases, aliases)

            # 1. the index the fund itself states
            bid = resolve_index(have, stated) if stated else None
            if bid:
                rows.append((code, bid, stated, "EXACT", 1.0, None))
                counts["EXACT"] += 1
                if nm_idx and index_ids[nm_idx] != bid:
                    # The statement wins, but two sources disagreeing is worth a human look.
                    counts["CONFLICT"] += 1
                    review.append((code, name, category, stated, nm_idx, "CONFLICT",
                                   "the fund states %s but its name says %s -- check which is right"
                                   % (stated, nm_idx)))
                continue

            # 1b. the index an index fund / ETF names in its own name
            if nm_idx:
                note = ("tracked index read from the fund's name (%r); %s"
                        % (nm_detail, "stated benchmark %s is not an index we hold" % stated if stated
                           else "its own stated benchmark was not available"))
                rows.append((code, index_ids[nm_idx], stated, "EXACT",
                             ALIAS_CONFIDENCE if nm_kind == "alias" else NAME_CONFIDENCE, note))
                counts["NAME"] += 1
                name_hits.setdefault(nm_idx, []).append(name)
                continue

            # 2. a near-equivalent of it
            if key and key in PROXIES:
                target, conf = PROXIES[key]
                bid = resolve_index(have, target)
                if bid:
                    note = "fund states %s; shown against %s" % (stated, target)
                    rows.append((code, bid, stated, "PROXY", conf, note))
                    counts["PROXY"] += 1
                    review.append((code, name, category, stated, target,
                                   "PROXY", "confirm the substitution is fair"))
                    continue

            # 3. the category's AMFI Tier-1 benchmark
            cat_key = norm(category)
            if (not args.no_category
                    and cat_key in CATEGORY_BENCHMARK
                    and cat_key not in CATEGORY_EXCLUDED):
                target, conf = CATEGORY_BENCHMARK[cat_key]
                conf = min(conf, CATEGORY_MAX_CONFIDENCE)
                bid = resolve_index(have, target)
                if bid:
                    note = ("category benchmark for %s; the scheme's own "
                            "stated benchmark was not available" % category)
                    rows.append((code, bid, stated, "CATEGORY", conf, note))
                    counts["CATEGORY"] += 1
                    cat_hits[category] = cat_hits.get(category, 0) + 1
                    review.append((code, name, category, stated or "", target,
                                   "CATEGORY",
                                   "category default -- confirm or override"))
                    continue

            # 4. nothing defensible
            if stated:
                counts["NONE_INDEX"] += 1
                rows.append((code, None, stated, "NONE", 0,
                             "index not loaded: %s" % stated))
                review.append((code, name, category, stated, "", "NONE",
                               "download %s, or leave unmapped" % stated))
            else:
                counts["NONE_NO_REF"] += 1
                rows.append((code, None, None, "NONE", 0,
                             "fund not present in the reference files"))
                if PASSIVE_MARKERS & set(name_tokens(name)):
                    # An index fund / ETF whose name we could not turn into an index we hold: say what
                    # it looks like it tracks. Counted across funds, this is the download list.
                    why = nm_kind if nm_idx is None else ""
                    review.append((code, name, category, "", "", "NONE",
                                   "index fund/ETF: name suggests '%s' (%s) -- download that index, "
                                   "or check the name" % (tracked_guess(name), why)))
                else:
                    review.append((code, name, category, "", "", "NONE",
                                   "not in reference files"))

        total_none = counts["NONE_INDEX"] + counts["NONE_NO_REF"]
        print("  EXACT     %4d   index we hold, as stated by the fund" % counts["EXACT"])
        print("  NAME      %4d   index fund / ETF: the index its own name says it tracks" % counts["NAME"])
        if counts["CONFLICT"]:
            print("  (CONFLICT %4d   stated benchmark and fund name disagree -- see %s)"
                  % (counts["CONFLICT"], REVIEW_CSV))
        print("  PROXY     %4d   near-equivalent substituted" % counts["PROXY"])
        print("  CATEGORY  %4d   category benchmark, no stated one found"
              % counts["CATEGORY"])
        print("  NONE      %4d   no comparison shown" % total_none)
        print("      %4d   fund found, but its index is not loaded"
              % counts["NONE_INDEX"])
        print("      %4d   fund not in any reference file we have"
              % counts["NONE_NO_REF"])
        if counts["SKIPPED"]:
            print("  SKIPPED   %4d   human-confirmed, left as-is" % counts["SKIPPED"])

        if name_hits:
            print("\nIndex funds / ETFs linked by name:")
            for idx, names in sorted(name_hits.items(), key=lambda kv: -len(kv[1])):
                print("  %-40s %4d" % (idx, len(names)))
                if args.list_name_matches:
                    for n in sorted(names):
                        print("        %s" % n)

        if cat_hits:
            print("\nCategory fallback applied to:")
            for cat, n in sorted(cat_hits.items(), key=lambda kv: -kv[1]):
                target, conf = CATEGORY_BENCHMARK[norm(cat)]
                print("  %-26s %4d  -> %-24s conf %.2f" % (cat, n, target, conf))

        # What this run would change against what is stored, so a re-run is never a leap of faith.
        with conn.cursor() as cur:
            cur.execute("SELECT scheme_code, benchmark_id, match_type FROM mf_benchmark_map")
            stored = {c: (b, t) for c, b, t in cur.fetchall()}
        name_of_id = {b: n for n, b in index_ids.items()}
        delta = {"same": 0, "gained": 0, "new_none": 0, "relabelled": 0, "switched": [], "lost": []}
        fund_name = {c: n for c, n, _ in funds}
        for code, new_bid, _stated, mtype, _conf, _note in rows:
            old = stored.get(code)
            if old is None:
                delta["gained" if new_bid else "new_none"] += 1
            elif old == (new_bid, mtype):
                delta["same"] += 1
            elif old[0] == new_bid:
                delta["relabelled"] += 1
            elif old[0] is None:
                delta["gained"] += 1
            elif new_bid is None:
                delta["lost"].append((fund_name[code], name_of_id.get(old[0])))
            else:
                delta["switched"].append((fund_name[code], name_of_id.get(old[0]), name_of_id.get(new_bid)))
        print("\nAgainst what is stored now: %d unchanged | %d GAIN a benchmark | %d new rows with none | "
              "%d same benchmark, different label | %d SWITCH benchmark | %d LOSE their benchmark"
              % (delta["same"], delta["gained"], delta["new_none"], delta["relabelled"],
                 len(delta["switched"]), len(delta["lost"])))
        for n, o, nw in delta["switched"][:15]:
            print("   SWITCH  %-60s %s -> %s" % (n[:60], o, nw))
        for n, o in delta["lost"][:15]:
            print("   LOSE    %-60s had %s" % (n[:60], o))

        with open(REVIEW_CSV, "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(["scheme_code", "scheme_name", "category",
                        "stated_benchmark", "shown_against", "match_type",
                        "action_needed"])
            w.writerows(review)
        print("\nWrote %s -- %d fund(s) needing a decision" % (REVIEW_CSV, len(review)))

        if args.dry_run:
            print("\nDRY RUN -- mf_benchmark_map not written.")
            return

        with conn.cursor() as cur:
            # reviewed rows were filtered out above, so nothing here can
            # overwrite a human decision.
            cur.executemany("""
                INSERT INTO mf_benchmark_map
                    (scheme_code, benchmark_id, stated_benchmark,
                     match_type, confidence, note, updated_at)
                VALUES (%s, %s, %s, %s, %s, %s, NOW())
                ON CONFLICT (scheme_code) DO UPDATE SET
                    benchmark_id     = EXCLUDED.benchmark_id,
                    stated_benchmark = EXCLUDED.stated_benchmark,
                    match_type       = EXCLUDED.match_type,
                    confidence       = EXCLUDED.confidence,
                    note             = EXCLUDED.note,
                    updated_at       = NOW()
                WHERE NOT mf_benchmark_map.reviewed
            """, rows)
            conn.commit()
        print("mf_benchmark_map written: %d row(s)" % len(rows))

        with conn.cursor() as cur:
            cur.execute("""
                SELECT COALESCE(b.display_name, '(none)'), map.match_type,
                       COUNT(*)
                FROM mf_benchmark_map map
                LEFT JOIN benchmark_master b USING (benchmark_id)
                GROUP BY 1, 2 ORDER BY 3 DESC
            """)
            print("\nMapping summary:")
            for display, match, n in cur.fetchall():
                print("  %-30s %-6s %4d" % (display, match, n))


if __name__ == "__main__":
    main()
