"""
parse_holdings.py -- generic AMC monthly-portfolio parser.

The five AMCs you downloaded use five different layouts. Rather than
writing five parsers, this finds the header row by LOOKING for it, then
maps columns by name. A sixth AMC will usually work with no changes.
"""

import re
import glob
import os
import shutil
import tempfile
from datetime import date

from openpyxl import load_workbook

# Some AMCs write "% to NAV"; HSBC writes "Percentage to Net Assets".
# One pattern covers both, plus "% of NAV" and "% to AUM".
PCT_HEADER_RE = r"(?:%|percentage|percent)\s*(?:to|of)\s*(?:nav|net asset|aum)"

# ---------------------------------------------------------------------
# Column detection: each field lists patterns that identify it.
# Order matters -- first match wins.
# ---------------------------------------------------------------------
COLUMN_PATTERNS = {
    "isin":     [r"^isin"],
    "name":     [r"name of the instrument", r"company/issuer", r"instrument\s*/\s*issuer",
                 r"name of instrument", r"^name$"],
    "industry": [r"industry", r"rating"],
    "quantity": [r"^quantity", r"^qty"],
    # UTI writes "MARKET-VALUE" with a hyphen, and the old pair of patterns
    # ("market/fair value", "market value") matched neither -- so the value
    # column never mapped and every market_value loaded NULL for all 18 of
    # its funds. Nothing complained: the sanity checks below test the percent
    # total only, so the files looked healthy with the column empty. One
    # pattern now covers hyphen, slash and space, with "fair" optional.
    "value":    [r"market\s*[-/]?\s*(?:fair\s*)?value", r"exposure/market",
                 r"amount", r"value.*lakh", r"value.*lac"],
    "pct":      [PCT_HEADER_RE],
}

# Section headings that tell us what KIND of instrument follows
SECTION_TYPES = [
    (r"equity\s*&?\s*equity related|equity shares|^equity$", "EQUITY"),
    (r"debt|bond|debenture|ncd|government securit|g-sec|gilt|treasury|t-bill", "DEBT"),
    (r"money market|certificate of deposit|commercial paper", "DEBT"),
    (r"treps|tri-?party|repo|reverse repo", "TREPS"),
    (r"cash|net receivable|net current asset|net assets", "CASH"),
    (r"reit|invit", "REIT"),
    (r"foreign|overseas|adr|gdr", "FOREIGN"),
    (r"mutual fund unit|units of", "OTHER"),
]

ISIN_RE = re.compile(r"^[A-Z]{2}[A-Z0-9]{9}\d$")

# Above this, the cell in the percentage column is not a percentage.
#
# Deliberately loose. Percentages arrive either as 9.86 or as 0.0986, so a
# value in the hundreds might still be a real figure read at the wrong
# scale and must survive to reach scale detection. Nothing legitimate is
# above a thousand. The tight "over 100%" check stays where it is, after
# scaling, and still reports.
PCT_ABSURD = 1000.0


def norm(text):
    """Squash whitespace and lowercase, so 'Market/Fair Value\n(Rs. in Lacs' matches."""
    if text is None:
        return ""
    return re.sub(r"\s+", " ", str(text)).strip().lower()


def find_header(rows, max_scan=30):
    """A header row is one that contains an ISIN column AND a % column."""
    for idx, row in enumerate(rows[:max_scan]):
        cells = [norm(c) for c in row]
        has_isin = any(re.match(r"^isin", c) for c in cells)
        has_pct = any(re.search(PCT_HEADER_RE, c) for c in cells)
        if has_isin and has_pct:
            return idx, cells
    return None, None


def map_columns(header_cells):
    mapping = {}
    for field, patterns in COLUMN_PATTERNS.items():
        for col_idx, cell in enumerate(header_cells):
            if col_idx in mapping.values():
                continue
            if any(re.search(p, cell) for p in patterns):
                mapping[field] = col_idx
                break
    return mapping


BOILERPLATE = ("portfolio", "statement", "as on", "registered office",
               "cin:", "investment manager", "asset management",
               "back to index", "disclaimer", "notes")


def looks_like_amc_name(text):
    """'SBI Mutual Fund' is the house, not the scheme. Short names ending
    in 'mutual fund' are the AMC."""
    words = text.split()
    return len(words) <= 4 and text.lower().rstrip().endswith("mutual fund")


def clean_fund_name(text):
    # Strip the parenthetical description FIRST -- some run 200+ chars,
    # and measuring length before stripping rejects valid names.
    return re.sub(r"\s+", " ", text.split("(")[0]).strip(" -:\t")


# 'portfolio of X', 'portfolio for X', and Zerodha's
# 'MONTHLY PORTFOLIO STATEMENT OF X FOR JULY 2026'. The trailing group
# eats a date suffix in any of its usual shapes so it doesn't end up
# glued to the scheme name.
TITLE_WRAPPER_RE = re.compile(
    r"portfolio\s+(?:statement\s+)?(?:of|for)\s+"
    r"(.+?)"
    r"(?:\s+(?:as\s+(?:on|at)|for\s+the\s+(?:month|period|quarter)|"
    r"for\s+\w+\s+\d{4}|as\s+of)\b.*)?$",
    re.I)


def unwrap_title(text):
    """'Portfolio of Kotak Nifty200 Value 30 Index Fund as on 31-Jul-2026'
    -> 'Kotak Nifty200 Value 30 Index Fund'.  Returns None if the line is
    not a title of that shape."""
    m = TITLE_WRAPPER_RE.search(re.sub(r"\s+", " ", str(text)).strip())
    if m:
        inner = m.group(1).strip(" -:,")
        # 'Portfolio for the month of July 2026' unwraps to a date phrase,
        # not a scheme. Same test step 2 applies: a scheme names itself.
        if not any(w in inner.lower() for w in ("fund", "scheme", "plan", "etf")):
            return None
        if 3 < len(inner) < 120:
            return inner
    return None


# An AMC's internal scheme code glued to the front of the name, where the
# code is NOT this sheet's name so strip_sheet_code cannot catch it: Groww
# writes "IB01-Groww Large Cap Fund" on a sheet called "BC". The name then
# normalises to ib01growwlargecapfund, matches no scheme, and the promote
# maps zero funds -- and "groww%" does not even match a string starting IB01.
#
# Requiring a DIGIT in the code is what makes this safe. Every AMC prefix is
# letters only, so "UTI - Mid Cap Fund" and "HDFC ELSS Tax saver" are
# untouched, while IB01, IB70 and the like are removed. "360 ONE" starts with
# digits rather than letters, so it does not match either.
CODE_PREFIX_RE = re.compile(r"^[A-Za-z]{1,4}\d{1,3}\s*[-:\u2013]\s*(?=\S)")


def strip_code_prefix(name):
    if not name:
        return name
    stripped = CODE_PREFIX_RE.sub("", name, count=1).strip()
    return stripped if len(stripped) > 3 else name


def strip_sheet_code(name, sheet_name):
    """'WCAR-THE WEALTH COMPANY ARBITRAGE FUND' -> 'THE WEALTH COMPANY
    ARBITRAGE FUND', but only when the prefix IS this sheet's name.

    The prefix is the AMC's internal scheme code, and the sheet is named
    after it, so requiring the two to agree is what makes this safe.
    Stripping any short prefix before a dash would turn a genuine AMFI
    name like 'UTI - Mid Cap Fund' into 'Mid Cap Fund'.
    """
    if not name or not sheet_name:
        return name
    code = re.escape(sheet_name.strip())
    stripped = re.sub(r"^%s\s*[-:\u2013]\s*" % code, "", name, count=1, flags=re.I)
    out = stripped.strip() if len(stripped.strip()) > 3 else name
    return strip_code_prefix(out)


def find_fund_name(rows, header_idx, sheet_name=None):
    """Look above the header row for the scheme name."""

    # 0. A wrapped title line. Checked FIRST because the wrapper words
    #    ("portfolio", "as on") are in BOILERPLATE and would otherwise
    #    cause the whole line -- name included -- to be discarded.
    for row in rows[:header_idx]:
        for cell in row:
            if cell is None:
                continue
            for line in str(cell).split("\n"):
                inner = unwrap_title(line)
                if inner and not looks_like_amc_name(inner):
                    return strip_sheet_code(clean_fund_name(inner), sheet_name)

    # 1. Explicit label, e.g. SBI's  'SCHEME NAME :' | 'SBI Contra Fund'
    for row in rows[:header_idx]:
        cells = list(row)
        for i, cell in enumerate(cells):
            if cell and re.search(r"scheme\s*name", str(cell), re.I):
                for follow in cells[i + 1:]:
                    if follow and len(str(follow).strip()) > 3:
                        return strip_sheet_code(clean_fund_name(str(follow)), sheet_name)

    # 2. Any line above the header that names a fund.
    #    A title block is often ONE cell holding several newline-separated
    #    lines ("HSBC Mutual Fund\nHSBC Flexi Cap Fund\n..."), so test each
    #    line on its own rather than the cell as a whole.
    for row in rows[:header_idx]:
        for cell in row:
            if cell is None:
                continue
            for line in str(cell).split("\n"):
                name = clean_fund_name(line)
                low = name.lower()
                if not (8 < len(name) < 120):
                    continue
                if not any(word in low for word in ("fund", "scheme", "plan", "etf")):
                    continue
                if any(bad in low for bad in BOILERPLATE):
                    continue
                if looks_like_amc_name(name):
                    continue
                return strip_sheet_code(name, sheet_name)
    return None


def fund_name_from_notes(wb):
    """Single-fund workbooks (HSBC) carry the full scheme name on a Notes
    sheet as 'Notes: <name>'. More reliable than the title block, which
    truncates names whose own text contains brackets."""
    for sheet_name in wb.sheetnames:
        if sheet_name.strip().lower() != "notes":
            continue
        ws = wb[sheet_name]
        for row in ws.iter_rows(min_row=1, max_row=3, values_only=True):
            for cell in row:
                if cell is None:
                    continue
                text = re.sub(r"\s+", " ", str(cell)).strip()
                if text.lower().startswith("notes:"):
                    text = text[6:].strip(" -:")
                    if 3 < len(text) < 120:
                        return text
    return None


def build_sheet_index(wb):
    """Some AMCs name each sheet with a code (V3I, A50) and put a
    code -> scheme-name lookup on a separate sheet. Find that lookup by
    testing whether a sheet's first column holds OTHER sheets' names."""
    sheet_names = {s.strip() for s in wb.sheetnames}
    best = {}
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        mapping = {}
        for row in ws.iter_rows(min_row=1, max_row=200, max_col=4, values_only=True):
            cells = [c for c in row if c is not None]
            if len(cells) < 2:
                continue
            code = str(cells[0]).strip()
            label = re.sub(r"\s+", " ", str(cells[1])).strip()
            if code in sheet_names and code != sheet_name and 3 < len(label) < 120:
                mapping[code] = label
        if len(mapping) > len(best):
            best = mapping
    # One or two accidental matches prove nothing; a real index covers most sheets.
    return best if len(best) >= max(3, len(sheet_names) // 4) else {}


def classify_section(text):
    low = norm(text)
    for pattern, kind in SECTION_TYPES:
        if re.search(pattern, low):
            return kind
    return None


def to_float(value):
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).replace(",", "").replace("%", "").strip()
    if text in ("", "-", "NA", "N.A."):
        return None
    try:
        return float(text)
    except ValueError:
        return None


def parse_sheet(rows, sheet_name, name_hint=None):
    header_idx, header_cells = find_header(rows)
    if header_idx is None:
        return None

    cols = map_columns(header_cells)
    if "isin" not in cols or "pct" not in cols:
        return None

    # A name from the sheet-index lookup skipped clean_fund_name, unlike
    # every other path -- so an AMC that writes "Nippon India Flexi Cap
    # Fund (An open ended dynamic equity scheme...)" on its index sheet
    # produced a name matching nothing in mf_scheme, and the promote
    # mapped zero funds.
    fund_name = ((clean_fund_name(name_hint) if name_hint else None)
                 or find_fund_name(rows, header_idx, sheet_name)
                 or sheet_name)
    holdings = []
    current_type = "EQUITY"
    net_assets_pct = None

    for row in rows[header_idx + 1:]:
        cells = list(row) + [None] * 12

        isin = str(cells[cols["isin"]]).strip() if cells[cols["isin"]] else ""

        name = ""
        if cols.get("name") is not None:
            if cells[cols["name"]]:
                name = str(cells[cols["name"]]).strip()
            else:
                # The header cell may be merged across several columns, so the
                # value sits to the RIGHT of where the header word was found.
                # Scan forward, stopping before the ISIN column.
                for probe in range(cols["name"] + 1, cols["isin"]):
                    v = cells[probe]
                    if v and str(v).strip():
                        name = str(v).strip()
                        break

        # A row with text but no valid ISIN is a section heading
        if not ISIN_RE.match(isin.upper()):
            heading = name or isin
            # The "Total Net Assets as on <date>" row states the whole in the
            # file's own units -- 1 if fractions, 100 if percents. Far more
            # reliable than summing holdings, which falls short whenever a
            # fund holds anything without an ISIN.
            if re.search(r"total net asset", norm(heading)):
                stated = to_float(cells[cols["pct"]])
                if stated is not None:
                    net_assets_pct = stated
            kind = classify_section(heading)
            if kind:
                current_type = kind
            continue

        holdings.append({
            "isin": isin.upper(),
            "name": name,
            "industry": str(cells[cols["industry"]]).strip() if cols.get("industry") is not None and cells[cols["industry"]] else None,
            "quantity": to_float(cells[cols["quantity"]]) if cols.get("quantity") is not None else None,
            "value": to_float(cells[cols["value"]]) if cols.get("value") is not None else None,
            "pct": to_float(cells[cols["pct"]]),
            "type": current_type,
        })

    if not holdings:
        return None

    # ---- ROWS WHOSE PERCENTAGE IS NOT A PERCENTAGE ----
    #
    # ABSL's August file prints rupee figures in the "% to Net Assets"
    # cell for defaulted IL&FS paper in three of its debt funds -- one
    # row read 1,044,498,811.69. That is not a mis-mapped column: the
    # value column on the same row holds a different number again. It is
    # what the AMC published.
    #
    # Quarantined BEFORE scale detection, and that ordering is the whole
    # point. `total` below decides whether every holding in the fund gets
    # multiplied by 100, and one billion-scale row makes that total
    # meaningless -- so a single bad cell would silently rescale every
    # OTHER holding in the same fund. Removing it first leaves the rest
    # summing to about 100, and the scale decision is made on real data.
    unreadable = []
    for h in holdings:
        if h["pct"] is not None and abs(h["pct"]) > PCT_ABSURD:
            unreadable.append((h["name"], h["pct"]))
            # None, not zero. "We could not read this" and "the fund holds
            # none of it" are different statements, and only one of them
            # is true. The row keeps its name, quantity and value.
            h["pct"] = None

    # ---- SCALE DETECTION ----
    # Some AMCs write 9.86 for 9.86%. Others write 0.0986. Prefer the
    # stated Total Net Assets row; fall back to summing the holdings.
    total = sum(h["pct"] for h in holdings if h["pct"] is not None)

    scale = 1.0
    basis = "none"
    if net_assets_pct is not None and 0.5 < net_assets_pct < 1.8:
        scale, basis = 100.0, "total-row"
    elif net_assets_pct is not None and 50 < net_assets_pct < 180:
        scale, basis = 1.0, "total-row"
    elif 0.5 < total < 1.8:
        scale, basis = 100.0, "sum"

    if scale != 1.0:
        for h in holdings:
            if h["pct"] is not None:
                h["pct"] *= scale

    # ---- SANITY CHECKS ----
    # Catch corruption that would otherwise load silently and skew scores.
    problems = []
    scaled_total = total * scale
    if not (50.0 <= scaled_total <= 150.0):
        problems.append("total %.2f%% outside 50-150" % scaled_total)
    over = [h["name"] for h in holdings if h["pct"] is not None and h["pct"] > 100.0]
    if over:
        problems.append("%d holding(s) above 100%%: %s" % (len(over), over[0][:40]))
    if unreadable:
        problems.append(
            "%d row(s) had no readable percentage and were loaded without one "
            "(worst: %s = %.0f)"
            % (len(unreadable), (unreadable[0][0] or "?")[:34], unreadable[0][1]))

    return {
        "fund_name": fund_name,
        "sheet": sheet_name,
        "holdings": holdings,
        "raw_total": total,
        "scale_applied": scale,
        "scale_basis": basis,
        "pct_total": scaled_total,
        "unreadable_pct": len(unreadable),
        "problems": problems,
    }


# ---------------------------------------------------------------------
# Genuinely old .xls files cannot be read by openpyxl, so pandas + xlrd
# reads them instead. But pandas exposes sheet_names and DataFrames,
# while every function below expects openpyxl's sheetnames / wb[name] /
# ws.iter_rows(). These two classes make a pandas workbook answer to the
# openpyxl calls, so the parser needs no special case for old files.
# ---------------------------------------------------------------------
class _XlsSheet:
    """One sheet, presented the way openpyxl presents one."""

    def __init__(self, frame):
        self._frame = frame

    @property
    def max_column(self):
        return self._frame.shape[1]

    def iter_rows(self, min_row=None, max_row=None, max_col=None,
                  values_only=True):
        import pandas as pd
        frame = self._frame
        if max_col:
            frame = frame.iloc[:, :max_col]
        start = (min_row - 1) if min_row else 0
        stop = max_row if max_row else None
        for row in frame.iloc[start:stop].itertuples(index=False, name=None):
            # pandas uses NaN for blanks; openpyxl uses None, and every
            # emptiness test below is written against None.
            yield tuple(None if (v is None or (isinstance(v, float) and pd.isna(v)))
                        else v for v in row)


class _XlsBook:
    def __init__(self, path):
        import pandas as pd
        self._xl = pd.ExcelFile(path, engine="xlrd")
        self.sheetnames = list(self._xl.sheet_names)

    def __getitem__(self, name):
        # header=None keeps the title block as data -- find_header() looks
        # for the header row itself rather than assuming pandas found it.
        return _XlsSheet(self._xl.parse(name, header=None, dtype=object))


def open_workbook_any(path):
    """Some AMCs ship an .xlsx but name it .xls. openpyxl trusts the
    extension, so copy to a correct one when needed."""
    if path.lower().endswith(".xls"):
        with open(path, "rb") as f:
            magic = f.read(4)
        if magic[:2] == b"PK":          # a zip => really xlsx
            tmp = os.path.join(tempfile.mkdtemp(), "fixed.xlsx")
            shutil.copy(path, tmp)
            return load_workbook(tmp, read_only=True, data_only=True)
        return _XlsBook(path)           # genuinely old format
    return load_workbook(path, read_only=True, data_only=True)


def parse_file(path):
    wb = open_workbook_any(path)
    sheet_index = build_sheet_index(wb) if hasattr(wb, "sheetnames") else {}
    results = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = [r for r in ws.iter_rows(values_only=True)]
        parsed = parse_sheet(rows, sheet_name, sheet_index.get(sheet_name.strip()))
        if parsed:
            parsed["source_file"] = os.path.basename(path)
            results.append(parsed)

    # One holdings sheet in the workbook means one fund per file. Only then
    # can a workbook-level Notes sheet safely name it -- in a multi-fund
    # workbook the same Notes name would be stamped onto every scheme.
    if len(results) == 1:
        notes_name = fund_name_from_notes(wb)
        if notes_name:
            results[0]["title_name"] = results[0]["fund_name"]
            results[0]["fund_name"] = notes_name

    return results


if __name__ == "__main__":
    import sys
    for pattern in sys.argv[1:]:
        for path in glob.glob(pattern):
            try:
                for r in parse_file(path):
                    total = sum(h["pct"] for h in r["holdings"] if h["pct"])
                    eq = sum(1 for h in r["holdings"] if h["type"] == "EQUITY")
                    flag = "  <-- CHECK: " + "; ".join(r["problems"]) if r.get("problems") else ""
                    print(f"{r['fund_name'][:48]:<48} | {len(r['holdings']):>4} rows "
                          f"| {eq:>4} eq | sum {total:>7.2f}% | x{r['scale_applied']:g}{flag}")
            except Exception as e:
                print(f"ERROR {os.path.basename(path)}: {e}")
